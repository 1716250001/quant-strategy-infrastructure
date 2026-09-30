# -*- coding: utf-8 -*-
"""
tools/export_lof.py — 导出场内 LOF 完整列表为 Excel
====================================================
重构要点：
  - 原实现把全流程写在模块顶层（import 即执行），现封装为
    build_lof_table() / export_excel() / main()，可被复用与测试
  - 输出路径原为硬编码绝对路径，现统一取 common.paths.DOC_DIR
"""
import os

import openpyxl
import pandas as pd

from common.paths import DOC_DIR, FUND_DAILY_DIR, META_DIR
from common import reader
from tools.utils import classify_fund


def invest_desc(row):
    """按名称推断 LOF 投资方向"""
    name = row["简称"]
    ftype = row["基金类型"]
    if "REIT" in name or "REITs" in name:
        if "产业园" in name or "产业园区" in name or "工业园" in name:
            return "产业园REITs"
        elif "高速" in name or "公路" in name:
            return "高速公路REITs"
        elif "保障房" in name or "保障性住房" in name or "租赁住房" in name or "住房" in name:
            return "保障房/租赁住房REITs"
        elif "能源" in name or "光伏" in name or "风电" in name or "发电" in name:
            return "新能源/能源REITs"
        elif "仓储" in name or "物流" in name:
            return "仓储物流REITs"
        elif "消费" in name or "商业" in name:
            return "消费/商业REITs"
        elif "水利" in name or "水务" in name:
            return "水务REITs"
        elif "污水处理" in name:
            return "环保REITs"
        elif "医药" in name or "医疗" in name:
            return "医药REITs"
        else:
            return "REITs"
    if "QDII" in name:
        return "QDII-LOF（海外投资）"
    if "FOF" in name:
        return "FOF-LOF（基金中基金）"
    if "债" in name or "债券" in name:
        return "债券型LOF"
    if "股票" in name:
        return "股票型LOF"
    if "混合" in name:
        return "混合型LOF"
    return ftype


def build_lof_table():
    """读取场内基金元数据，返回按代码排序的 LOF 明细 DataFrame"""
    fb = pd.read_parquet(os.path.join(META_DIR, "fund_basic.parquet"))
    fb["delist_date"] = fb["delist_date"].fillna("")
    # fund_daily 已改为按年分区，无法再从文件名取代码（旧写法在 by_year
    # 下会把 "2026" 当成代码），改走 reader.codes_from_store。
    fund_files = set(reader.codes_from_store("fund_daily"))
    on_market = fb[fb["ts_code"].isin(fund_files)].copy()

    on_market["code"] = (on_market["ts_code"]
                         .str.replace(".SH", "").str.replace(".SZ", "").str.replace(".BJ", ""))
    on_market["suffix"] = on_market["ts_code"].str[-2:]
    on_market["market_class"] = on_market.apply(classify_fund, axis=1)
    lof = on_market[on_market["market_class"] == "LOF"].copy()

    lof["上市日期"] = lof["list_date"].fillna("")
    lof = lof.rename(columns={"ts_code": "代码", "name": "简称", "fund_type": "基金类型"})
    lof = lof[["代码", "简称", "基金类型", "上市日期"]].copy()
    lof = lof.sort_values("代码").reset_index(drop=True)
    lof.insert(0, "序号", range(1, len(lof) + 1))
    lof["投资方向"] = lof.apply(invest_desc, axis=1)
    return lof


def export_excel(lof, out_path):
    """写出带样式的 Excel 文件"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "场内LOF列表"

    headers = ["序号", "代码", "简称", "基金类型", "投资方向", "上市日期"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = openpyxl.styles.Font(bold=True, size=11, color="FFFFFF")
        cell.alignment = openpyxl.styles.Alignment(horizontal="center")
        cell.fill = openpyxl.styles.PatternFill(start_color="4472C4", end_color="4472C4",
                                                fill_type="solid")

    for _, row in lof.iterrows():
        ws.append([row["序号"], row["代码"], row["简称"], row["基金类型"],
                   row["投资方向"], row["上市日期"]])

    for col, width in zip("ABCDEF", [6, 14, 42, 10, 24, 12]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(lof) + 1}"
    wb.save(out_path)


def run(out_path=None):
    """导出场内 LOF 完整列表为 Excel，返回输出路径。

    供本文件 CLI 与统一入口 main.py 的 lof-export 子命令共用。
    """
    lof = build_lof_table()
    out_path = out_path or os.path.join(DOC_DIR, "场内LOF完整列表.xlsx")
    export_excel(lof, out_path)
    print(f"已保存: {out_path}")
    print(f"共 {len(lof)} 只场内LOF")
    print("\n投资方向分布:")
    print(lof["投资方向"].value_counts().to_string())
    return out_path


def main():
    run()


if __name__ == "__main__":
    main()
